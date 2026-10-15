import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

def export_executive_excel(csv_input, excel_output):
    print(f"📊 Generando informe Excel formateado desde {csv_input}...")
    
    try:
        df = pd.read_csv(csv_input)
    except FileNotFoundError:
        print(f"❌ Error: No se encontró el archivo '{csv_input}'.")
        return
    
    # Crear Workbook de openpyxl
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte Overprice Carrefour"
    ws.views.sheetView[0].showGridLines = True

    # 1. Título del Reporte
    ws.merge_cells("A1:G1")
    title_cell = ws["A1"]
    title_cell.value = "RETAIL PULSE — INFORME DE SOBREPRECIOS: CARREFOUR"
    title_cell.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40

    # Subtítulo con fecha
    ws.merge_cells("A2:G2")
    sub_cell = ws["A2"]
    sub_cell.value = f"Fecha de emisión: {datetime.now().strftime('%d/%m/%Y')} | Matriz de comparación por EAN exacto"
    sub_cell.font = Font(name="Calibri", size=10, italic=True, color="595959")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")

    # 2. Renombrar columnas para la presentación
    column_mapping = {
        "ean": "EAN",
        "product_name": "Producto",
        "Carrefour": "Carrefour ($)",
        "Coto": "Coto ($)",
        "Jumbo": "Jumbo ($)",
        "diff_abs": "Brecha ($)",
        "diff_pct": "Brecha (%)"
    }
    
    # Asegurar orden de columnas
    cols = [c for c in ["ean", "product_name", "Carrefour", "Coto", "Jumbo", "diff_abs", "diff_pct"] if c in df.columns]
    df_display = df[cols].rename(columns=column_mapping)

    # 3. Encabezados de la tabla
    start_row = 4
    headers = list(df_display.columns)
    
    header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_num)
        cell.value = header_title
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[start_row].height = 25

    # 4. Volcar datos y formatear celdas
    red_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") # Alerta roja suave
    orange_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # Alerta amarilla/naranja

    for row_idx, row_data in enumerate(df_display.values, start_row + 1):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = value
            cell.border = thin_border
            col_name = headers[col_idx - 1]

            # Alineación y formatos de números
            if col_name == "EAN":
                cell.alignment = Alignment(horizontal="center")
                cell.number_format = "@"
            elif col_name == "Producto":
                cell.alignment = Alignment(horizontal="left")
            elif "($)" in col_name:
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = '$ #,##0.00'
            elif "(%)" in col_name:
                cell.alignment = Alignment(horizontal="right")
                cell.number_format = '0.00"%"'
                
                # Resaltado condicional por porcentaje de sobreprecio
                if pd.notna(value):
                    if value >= 30:
                        cell.fill = red_fill
                        cell.font = Font(bold=True, color="C00000")
                    elif value >= 10:
                        cell.fill = orange_fill
                        cell.font = Font(bold=True, color="B25900")

    # 5. Autoajustar ancho de columnas de forma segura
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for row_idx in range(start_row, ws.max_row + 1):
            cell_val = ws.cell(row=row_idx, column=col_idx).value
            if cell_val is not None:
                max_len = max(max_len, len(str(cell_val)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
    
    # Ancho fijo especial para Producto
    ws.column_dimensions['B'].width = 50

    wb.save(excel_output)
    print(f"✅ ¡Informe Excel profesional generado con éxito!: {excel_output}")

if __name__ == "__main__":
    input_csv = "carrefour_reporte_precios_altos_20261013.csv"
    output_xlsx = "Informe_Ejecutivo_Carrefour_Overprice.xlsx"
    export_executive_excel(input_csv, output_xlsx)